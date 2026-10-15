"""Turn labels/to_label.csv into labels/gold.xlsx with dropdowns, so labelling is click-only.
Allowed values match extract.py's PROMPT. evaluate.py reads gold.xlsx directly.
Refuses to overwrite an existing gold.xlsx (it holds hand labels).
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE / "labels/gold.xlsx"
CHOICES = {
    "is_decision": ["TRUE", "FALSE"],
    "request_type": ["clinical_trial", "bioequivalence", "marketing_new_drug", "new_indication_or_change",
                     "protocol_amendment", "post_marketing", "other"],
    "local_ct_waiver_requested": ["TRUE", "FALSE"],
    "decision": ["approved", "approved_with_conditions", "rejected", "more_info_requested", "deferred", "other"],
    "reasons_stated": ["TRUE", "FALSE"],
}
assert not OUT.exists(), f"{OUT} already exists; delete it yourself if you really want a fresh sheet"

rows = list(csv.DictReader(open(HERE / "labels/to_label.csv", encoding="utf-8-sig")))
wb = Workbook(); ws = wb.active; ws.title = "label"
cols = ["key", "year", "committee", "text"] + list(CHOICES)
ws.append(cols)
for r in rows:
    ws.append([r["key"], r["year"], r["committee"], r["text"], "TRUE", "", "", "", ""])  # most rows are real decisions

widths = {"A": 10, "B": 6, "C": 28, "D": 90, "E": 12, "F": 24, "G": 14, "H": 24, "I": 12}
for c, w in widths.items(): ws.column_dimensions[c].width = w
for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDEBF7")
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "E2"
for i, (field, opts) in enumerate(CHOICES.items()):
    dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
    col = chr(ord("E") + i)
    dv.add(f"{col}2:{col}{len(rows) + 1}"); ws.add_data_validation(dv)

guide = wb.create_sheet("how to label")
for line in [
    ["field", "pick", "example"],
    ["is_decision", "FALSE only if the row is not an application (summary table, note, heading)", "'Total cases: 12' -> FALSE"],
    ["request_type", "what the firm asked for", ""],
    ["", "clinical_trial", "permission for a CT / CT protocol, any phase"],
    ["", "bioequivalence", "BE study"],
    ["", "marketing_new_drug", "import / manufacture and market a new drug"],
    ["", "new_indication_or_change", "new indication, dosage form, strength, route, age group"],
    ["", "protocol_amendment", "amendment to an approved CT protocol"],
    ["", "post_marketing", "PMS, Phase IV, PSUR"],
    ["local_ct_waiver_requested", "TRUE if the firm asked to skip the local (Indian) clinical trial", "'...by waiver of local clinical trial' -> TRUE"],
    ["decision", "what the committee decided", ""],
    ["", "approved", "'recommended for grant of permission' (no conditions)"],
    ["", "approved_with_conditions", "'recommended ... subject to Phase IV / PMS / changes'"],
    ["", "rejected", "'not recommended', 'did not agree'"],
    ["", "more_info_requested", "'firm should submit revised protocol / data and present again'"],
    ["", "deferred", "'firm did not turn up', 'not ready for presentation'"],
    ["reasons_stated", "TRUE if the committee gives any why (efficacy, safety, unmet need, data gaps)", "'in view of unmet need...' -> TRUE"],
]: guide.append(line)
for c, w in {"A": 26, "B": 70, "C": 60}.items(): guide.column_dimensions[c].width = w
wb.save(OUT)
print(f"{len(rows)} rows -> {OUT}")
