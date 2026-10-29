"""Pick 40 of the 200 Claude-labelled items for Arav to check -> labels/human_verified.xlsx.
He fixes any wrong label in place and sets checked=YES; evaluate.py only scores checked rows.
"""
import random
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE / "labels/human_verified.xlsx"
assert not OUT.exists(), f"{OUT} exists (it may hold your checks); delete it yourself to rebuild"

wb = load_workbook(HERE / "labels/claude_labels.xlsx"); ws = wb["label"]
random.seed(40)
keep = set(random.sample(range(2, ws.max_row + 1), 40))
for r in range(ws.max_row, 1, -1):  # delete bottom-up so row numbers stay valid
    if r not in keep: ws.delete_rows(r)
ws.cell(1, 11, "checked")
dv = DataValidation(type="list", formula1='"YES"', allow_blank=True); ws.add_data_validation(dv)
dv.add("K2:K41")
for r in range(2, 42): ws.cell(r, 11).fill = PatternFill("solid", fgColor="FFF2CC")
ws.column_dimensions["K"].width = 10
wb.save(OUT)
print(f"40 rows -> {OUT}")
